"""Primitivas de desenho em PPTX, em pixels (1080x1350)."""
from __future__ import annotations

import random
from dataclasses import dataclass
from pathlib import Path

import numpy as np
from PIL import Image, ImageFilter, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from .markup import parse

EMU_PER_PX = 9525
ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def px(v: float) -> Emu:
    return Emu(int(round(v * EMU_PER_PX)))


def hex_rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


@dataclass
class TextStyle:
    family: str
    size: float                 # px
    line: float                 # px (entrelinha)
    color: str = "#FFFFFF"
    bold: bool = False
    italic: bool = False
    caps: bool = False
    tracking: float = 0.0       # em (0.12 = 12% do corpo)
    accent_family: str | None = None   # família usada em *ênfase*
    accent_italic: bool = True


def new_presentation(width: int, height: int) -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(width), px(height)
    return prs


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _line_spacing(paragraph, line_px: float) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    for old in pPr.findall(qn("a:lnSpc")):
        pPr.remove(old)
    ln = pPr.makeelement(qn("a:lnSpc"), {})
    ln.append(ln.makeelement(qn("a:spcPts"), {"val": str(int(round(line_px * 0.75 * 100)))}))
    pPr.insert(0, ln)


def _space_before(paragraph, px_value: float) -> None:
    if px_value <= 0:
        return
    pPr = paragraph._p.get_or_add_pPr()
    sb = pPr.makeelement(qn("a:spcBef"), {})
    sb.append(sb.makeelement(qn("a:spcPts"), {"val": str(int(round(px_value * 0.75 * 100)))}))
    pPr.append(sb)


def text_box(slide, name: str, box: tuple[float, float, float, float], paragraphs: list[str],
             style: TextStyle, align: str = "left", anchor: str = "top", paragraph_gap: float = 0.0):
    """paragraphs: lista de strings com marcação (ver markup.py); '\\n' quebra linha."""
    left, top, width, height = box
    shape = slide.shapes.add_textbox(px(left), px(top), px(width), px(height))
    shape.name = name
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = ANCHOR[anchor]
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for block in paragraphs:
        for j, line in enumerate(block.split("\n")):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            if not first and j == 0:
                _space_before(p, paragraph_gap)
            first = False
            p.alignment = ALIGN[align]
            _line_spacing(p, style.line)
            for span in parse(line):
                content = span.text.upper() if style.caps else span.text
                r = p.add_run()
                r.text = content
                f = r.font
                accent = span.accent and style.accent_family
                f.name = style.accent_family if accent else style.family
                f.italic = style.accent_italic if accent else style.italic
                f.bold = span.bold or style.bold
                f.underline = span.underline
                f.size = Pt(style.size * 0.75)
                f.color.rgb = hex_rgb(style.color)
                if style.tracking:
                    r._r.get_or_add_rPr().set("spc", str(int(round(style.tracking * style.size * 0.75 * 100))))
    return shape


def rect(slide, name, box, fill: str | None = None, outline: str | None = None, outline_px: float = 1.5,
         radius: float | None = None):
    left, top, width, height = box
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, px(left), px(top), px(width), px(height))
    shape.name = name
    if radius is not None:
        shape.adjustments[0] = min(0.5, radius / min(width, height))
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = hex_rgb(fill)
    else:
        shape.fill.background()
    if outline:
        shape.line.color.rgb = hex_rgb(outline)
        shape.line.width = Pt(outline_px * 0.75)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def hline(slide, name, x1, y, x2, color="#FFFFFF", weight_px=1.5):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x1), px(y), px(x2), px(y))
    line.name = name
    line.line.color.rgb = hex_rgb(color)
    line.line.width = Pt(weight_px * 0.75)
    return line


def segment(slide, name, x1, y1, x2, y2, color="#FFFFFF", weight_px=1.5):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x1), px(y1), px(x2), px(y2))
    line.name = name
    line.line.color.rgb = hex_rgb(color)
    line.line.width = Pt(weight_px * 0.75)
    return line


def cover_crop(img_w: int, img_h: int, box_w: float, box_h: float,
               focus: tuple[float, float] = (0.5, 0.5), zoom: float = 1.0) -> tuple[float, float, float, float]:
    """Frações (left, top, right, bottom) a cortar para a imagem cobrir a caixa.

    focus = ponto da foto (0..1) que fica no centro quando há sobra; zoom > 1 aproxima.
    """
    box_ratio, img_ratio = box_w / box_h, img_w / img_h
    vw, vh = (box_ratio / img_ratio, 1.0) if img_ratio > box_ratio else (1.0, img_ratio / box_ratio)
    zoom = max(zoom, 1.0)
    vw, vh = vw / zoom, vh / zoom
    fx, fy = focus
    sx = min(max(fx - vw / 2, 0.0), 1 - vw)
    sy = min(max(fy - vh / 2, 0.0), 1 - vh)
    return sx, sy, 1 - vw - sx, 1 - vh - sy


def picture(slide, name, path: Path, box, focus: tuple[float, float] = (0.5, 0.5), zoom: float = 1.0):
    """Insere a imagem cobrindo a caixa (object-fit: cover) com ponto de foco (0..1) e zoom."""
    left, top, width, height = box
    pic = slide.shapes.add_picture(str(path), px(left), px(top), px(width), px(height))
    pic.name = name
    with Image.open(path) as im:
        iw, ih = im.size
    l, t, r, b = cover_crop(iw, ih, width, height, focus, zoom)
    pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = l, t, r, b
    return pic


def region_luminance(path: Path, slot, region, focus=(0.5, 0.5), zoom: float = 1.0) -> tuple[float, float, float]:
    """(p20, média, p80) da luminância (0..1) da parte da foto que fica sob `region`.

    slot = caixa (left, top, w, h) onde a foto foi posta; region = caixa do texto,
    nas mesmas coordenadas do slide.
    """
    sl, st, sw, sh = slot
    rl, rt, rw, rh = region
    with Image.open(path) as im:
        im = im.convert("L")
        im.draft("L", (int(sw), int(sh)))
        l, t, r, b = cover_crop(im.width, im.height, sw, sh, focus, zoom)
        crop = im.crop((l * im.width, t * im.height, (1 - r) * im.width, (1 - b) * im.height))
        crop = crop.resize((max(1, int(sw / 4)), max(1, int(sh / 4))))
        x0 = max(0, int((rl - sl) / 4)); y0 = max(0, int((rt - st) / 4))
        x1 = min(crop.width, int((rl - sl + rw) / 4) + 1); y1 = min(crop.height, int((rt - st + rh) / 4) + 1)
        if x1 <= x0 or y1 <= y0:
            return 0.0, 0.0, 0.0
        values = np.sort(np.asarray(crop.crop((x0, y0, x1, y1)), dtype=np.float32).ravel()) / 255
        return float(values[int(len(values) * 0.2)]), float(values.mean()), float(values[int(len(values) * 0.8)])


def logo(slide, name, path: Path, left, top, width):
    with Image.open(path) as im:
        iw, ih = im.size
    pic = slide.shapes.add_picture(str(path), px(left), px(top), px(width), px(width * ih / iw))
    pic.name = name
    return pic


def prepare_photo(
    src: Path,
    dst: Path,
    max_side: int = 2160,
    quality: int = 88,
    neutralize: bool = False,
) -> Path:
    """Reduz a foto e, para banco externo, aplica o duotone neutro da marca."""
    dst.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(src) as im:
        im = im.convert("RGB")
        scale = max_side / max(im.size)
        if scale < 1:
            im = im.resize((round(im.width * scale), round(im.height * scale)), Image.LANCZOS)
        if neutralize:
            im = ImageOps.colorize(
                ImageOps.grayscale(im),
                black="#111111",
                mid="#777570",
                white="#E8E3DD",
            )
        im.save(dst, "JPEG", quality=quality)
    return dst


def _edge_fill(band: Image.Image, size: tuple[int, int], blur: float, rng) -> Image.Image:
    """Estica uma faixa de borda (1 px) até `size`, com desfoque e grão leve."""
    filled = np.asarray(band.resize(size, Image.NEAREST).filter(ImageFilter.GaussianBlur(blur)), dtype=np.int16)
    grain = rng.normal(0, 2.5, (size[1], size[0], 1)).astype(np.int16)
    return Image.fromarray(np.clip(filled + grain, 0, 255).astype(np.uint8))


def extend_canvas(src: Path, dst: Path, top: float = 0, left: float = 0, right: float = 0,
                  seed: int = 11) -> Path:
    """Estende o fundo de estúdio (frações da altura/largura), repetindo a cor da borda.

    Serve para abrir espaço de texto e afastar/diminuir o modelo sem cortar o corpo,
    como nos posts de estúdio publicados. Só faz sentido em fotos de fundo liso."""
    dst.parent.mkdir(parents=True, exist_ok=True)
    rng = np.random.default_rng(seed)
    with Image.open(src) as im:
        im = im.convert("RGB")
    w, h = im.size
    blur = max(w, h) / 80
    pl, pr = int(round(w * left)), int(round(w * right))
    if pl or pr:
        wide = Image.new("RGB", (w + pl + pr, h))
        wide.paste(im, (pl, 0))
        edge = max(4, w // 60)
        if pl:
            band = im.crop((0, 0, edge, h)).resize((1, h), Image.BOX)
            wide.paste(_edge_fill(band, (pl, h), blur, rng), (0, 0))
        if pr:
            band = im.crop((w - edge, 0, w, h)).resize((1, h), Image.BOX)
            wide.paste(_edge_fill(band, (pr, h), blur, rng), (pl + w, 0))
        im, w = wide, wide.width
    pt = int(round(h * top))
    if pt:
        tall = Image.new("RGB", (w, h + pt))
        band = im.crop((0, 0, w, max(4, h // 60))).resize((w, 1), Image.BOX)
        tall.paste(_edge_fill(band, (w, pt), blur, rng), (0, 0))
        tall.paste(im, (0, pt))
        im = tall
    seam = max(2, int(h * 0.01))
    if pt:
        im.paste(im.crop((0, pt - seam, w, pt + seam)).filter(ImageFilter.GaussianBlur(seam / 3)), (0, pt - seam))
    for x in [pl] if pl else []:
        im.paste(im.crop((x - seam, 0, x + seam, im.height)).filter(ImageFilter.GaussianBlur(seam / 3)), (x - seam, 0))
    if pr:
        x = im.width - pr
        im.paste(im.crop((x - seam, 0, x + seam, im.height)).filter(ImageFilter.GaussianBlur(seam / 3)), (x - seam, 0))
    im.save(dst, "JPEG", quality=88)
    return dst


def textured_black(dst: Path, width: int, height: int, base: str = "#0B0B0B", seed: int = 7) -> Path:
    """Fundo preto com grão sutil (procedural), como nas telas de manifesto da marca."""
    if dst.exists():
        return dst
    dst.parent.mkdir(parents=True, exist_ok=True)
    rnd = random.Random(seed)
    b = int(base.lstrip("#")[0:2], 16)
    small = Image.new("L", (width // 2, height // 2))
    small.putdata([max(0, min(255, b + int(rnd.gauss(0, 5)))) for _ in range((width // 2) * (height // 2))])
    im = small.resize((width, height), Image.BILINEAR).filter(ImageFilter.GaussianBlur(0.6)).convert("RGB")
    im.save(dst, "JPEG", quality=90)
    return dst
