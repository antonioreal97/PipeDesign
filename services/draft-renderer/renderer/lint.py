"""Confere o PPTX gerado contra as regras de importação e da marca."""
from __future__ import annotations

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from .brand import Brand
from .primitives import EMU_PER_PX


def lint(path, brand: Brand) -> list[str]:
    problems: list[str] = []
    prs = Presentation(str(path))
    if (prs.slide_width, prs.slide_height) != (brand.width * EMU_PER_PX, brand.height * EMU_PER_PX):
        problems.append("tamanho do slide diferente do canvas da marca")
    fonts = brand.data["fonts"]
    allowed = {fonts["display"]["family"], fonts["sans"]["family"]}
    palette = {c.upper().lstrip("#") for c in brand.data["colors"].values() if isinstance(c, str)}
    for n, slide in enumerate(prs.slides, 1):
        families = set()
        for shape in slide.shapes:
            where = f"slide {n} '{shape.name}'"
            if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                full_canvas = (
                    shape.left == 0 and shape.top == 0 and
                    shape.width == prs.slide_width and shape.height == prs.slide_height
                )
                if shape.fill.type is not None and shape.fill.type != 5 and not full_canvas:   # 5 = sem preenchimento
                    problems.append(f"{where}: forma preenchida (o padrão usa só contorno)")
            if not shape.has_text_frame:
                continue
            for p in shape.text_frame.paragraphs:
                if not p.runs:
                    continue
                pPr = p._p.pPr
                if pPr is None or pPr.find(qn("a:lnSpc")) is None:
                    problems.append(f"{where}: parágrafo sem entrelinha explícita")
                for r in p.runs:
                    fam = r.font.name
                    families.add(fam)
                    if fam not in allowed:
                        problems.append(f"{where}: fonte '{fam}' fora da tipografia editorial")
                    if "semibold" in (fam or "").lower():
                        problems.append(f"{where}: peso SemiBold não é reconhecido pelo Canva")
                    rgb = str(r.font.color.rgb) if r.font.color and r.font.color.type is not None else None
                    if rgb and rgb.upper() not in palette:
                        problems.append(f"{where}: cor #{rgb} fora da paleta")
        if len(families) > 2:
            problems.append(f"slide {n}: {len(families)} famílias ({', '.join(sorted(families))}); máximo 2")
    return problems
