"""Testes do draft-renderer. Rodar a partir de services/draft-renderer:

    python -m pytest -q
"""
from __future__ import annotations

import copy
import json
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

from renderer.brand import load_brand
from renderer.lint import lint
from renderer.markup import Span, parse, plain
from renderer.plan import validate
from renderer.primitives import EMU_PER_PX, cover_crop, extend_canvas, prepare_photo
from renderer.render import render

REPO = Path(__file__).resolve().parents[3]
BRANDS = REPO / "brands"


@pytest.fixture()
def client(tmp_path: Path) -> Path:
    """Pasta de cliente falsa: fotos de estúdio sintéticas + ícone."""
    root = tmp_path / "client"
    for rel, color in [("Fotos/a.jpg", (205, 205, 205)), ("Fotos/b.jpg", (40, 40, 40))]:
        p = root / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (600, 750), color).save(p)
    logo = root / "Logos/ICONE BRANCO SEM FUNDO.png"
    logo.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGBA", (100, 100), (255, 255, 255, 255)).save(logo)
    return root


@pytest.fixture()
def brand(client: Path):
    return load_brand(BRANDS, "exemplo", client)


def plan_all_layouts() -> dict:
    return {
        "brand": "exemplo",
        "slides": [
            {"layout": "capa_manchete", "photo": "Fotos/a.jpg", "kicker": "Por que o",
             "headline": "Tecido\nimporta?", "box": "Porque a peça acompanha você."},
            {"layout": "foto_nota", "photo": "Fotos/b.jpg", "position": "top_left",
             "paragraphs": ["Um **tecido agradável** muda a sensação."], "phrase": "E começa a pensar"},
            {"layout": "manifesto_preto", "title": "A elasticidade",
             "paragraphs": ["permite que o tecido __acompanhe__ o corpo.", "**Alongar. Correr.**"]},
            {"layout": "manifesto_frase", "text": "Elegância em movimento."},
            {"layout": "moodboard_cena", "side": "left",
             "scene": {"photo": "Fotos/b.jpg"}, "product": {"photo": "Fotos/a.jpg"},
             "detail": {"photo": "Fotos/a.jpg", "zoom": 2}, "time": "8:00",
             "lines": ["Daily com a equipe", "Responder e-mails"]},
            {"layout": "beneficio_duplo",
             "top": {"photo": "Fotos/a.jpg", "title": "Conforto\ntérmico", "text": "Regulação inteligente."},
             "bottom": {"photo": "Fotos/b.jpg", "title": "4-way stretch", "text": "Elasticidade nos quatro sentidos."}},
            {"layout": "manifesto_assinatura", "lines": ["**Tecnologia que você veste.**"],
             "tagline": "Elegância em movimento."},
        ],
    }


# ------------------------------------------------------------------ markup
def test_markup_spans():
    assert parse("Um **tecido** *leve* e __fino__.") == [
        Span("Um "), Span("tecido", bold=True), Span(" "), Span("leve", accent=True),
        Span(" e "), Span("fino", underline=True), Span("."),
    ]
    assert plain("**a** *b* __c__") == "a b c"


def test_markup_keeps_lone_asterisk():
    assert plain("10% * 2") == "10% * 2"


# ------------------------------------------------------------------ geometry
def test_cover_crop_same_ratio_is_identity():
    assert cover_crop(3000, 3750, 1080, 1350) == (0.0, 0.0, 0.0, 0.0)


def test_cover_crop_wide_slot_uses_focus():
    l, t, r, b = cover_crop(800, 1000, 1080, 675, focus=(0.5, 0.0))
    assert (l, r) == (0.0, 0.0) and t == 0.0 and b == pytest.approx(0.5)


def test_cover_crop_zoom_clamps_inside_image():
    l, t, r, b = cover_crop(1000, 1000, 500, 500, focus=(0.0, 1.0), zoom=2)
    assert l == 0.0 and b == pytest.approx(0.0) and r == pytest.approx(0.5) and t == pytest.approx(0.5)


def test_extend_canvas_sizes(tmp_path):
    src = tmp_path / "s.jpg"
    Image.new("RGB", (100, 200), (200, 200, 200)).save(src)
    out = extend_canvas(src, tmp_path / "o.jpg", top=0.5, left=0.2, right=0.3)
    with Image.open(out) as im:
        assert im.size == (150, 300)
        r, g, b = im.getpixel((5, 5))
        assert abs(r - 200) < 12


def test_prepare_photo_neutraliza_banco_externo(tmp_path):
    src = tmp_path / "color.jpg"
    Image.new("RGB", (40, 60), (240, 40, 20)).save(src)
    out = prepare_photo(src, tmp_path / "neutral.jpg", neutralize=True)
    with Image.open(out) as im:
        r, g, b = im.getpixel((20, 30))
    assert max(r, g, b) - min(r, g, b) < 12


# ------------------------------------------------------------------ plan
def test_valid_plan_has_no_errors(brand):
    rep = validate(plan_all_layouts(), brand)
    assert rep.errors == []


@pytest.mark.parametrize("mutate, message", [
    (lambda p: p["slides"].append({"layout": "banner"}), "layout desconhecido"),
    (lambda p: p["slides"][0].pop("headline"), "falta 'headline'"),
    (lambda p: p["slides"][0].update(headline="Uma\nduas\ntrês"), "máximo 2"),
    (lambda p: p["slides"][1].update(photo="Fotos/nao-existe.jpg"), "foto não encontrada"),
    (lambda p: p["slides"][1].update(position="centro"), "position"),
    (lambda p: p["slides"][2].update(title="Leve ⚡️"), "emoji"),
    (lambda p: p["slides"][0].update(extend={"bottom": 0.2}), "extend"),
    (lambda p: p.update(brand="outra"), "marca"),
])
def test_plan_errors(brand, mutate, message):
    plan = copy.deepcopy(plan_all_layouts())
    mutate(plan)
    rep = validate(plan, brand)
    assert any(message in e for e in rep.errors), rep.errors


def test_long_texts_are_warnings(brand):
    plan = plan_all_layouts()
    plan["slides"][1]["paragraphs"] = ["palavra " * 40]
    rep = validate(plan, brand)
    assert rep.ok and any("caracteres" in w for w in rep.warnings)


# ------------------------------------------------------------------ render
def test_render_all_layouts(brand, tmp_path):
    out, warnings = render(plan_all_layouts(), brand, tmp_path / "out.pptx", tmp_path / "cache")
    prs = Presentation(str(out))
    assert len(prs.slides) == 7
    assert (prs.slide_width, prs.slide_height) == (1080 * EMU_PER_PX, 1350 * EMU_PER_PX)
    assert lint(out, brand) == []

    fonts = brand.data["fonts"]
    runs = {}
    for n, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    if p.runs:
                        assert p._p.pPr.find(qn("a:lnSpc")) is not None
                    for r in p.runs:
                        runs[(n, shape.name, r.text)] = r
    # capa: manchete em caixa alta, serif itálica; texto escuro sobre foto clara
    head = runs[(1, "manchete", "TECIDO")]
    assert head.font.name == fonts["display"]["family"] and head.font.italic
    assert str(head.font.color.rgb) == brand.color("texto_escuro").lstrip("#")
    # foto escura -> texto branco; negrito preservado
    bold = runs[(2, "nota", "tecido agradável")]
    assert bold.font.bold and str(bold.font.color.rgb) == "FFFFFF"
    # sublinhado nas telas pretas
    assert runs[(3, "texto", "acompanhe")].font.underline
    # kicker com espaçamento entre letras
    kicker = runs[(1, "kicker", "POR QUE O")]
    assert int(kicker._r.rPr.get("spc")) > 0


def test_capa_box_is_outline_only(brand, tmp_path):
    out, _ = render(plan_all_layouts(), brand, tmp_path / "out.pptx", tmp_path / "cache")
    slide = Presentation(str(out)).slides[0]
    box = next(s for s in slide.shapes if s.name == "caixa")
    assert box.fill.type == 5          # sem preenchimento
    assert box.line.width > 0


def test_mixed_background_warns(brand, client, tmp_path):
    half = client / "Fotos/half.jpg"
    im = Image.new("RGB", (600, 750), (230, 230, 230))
    im.paste((20, 20, 20), (0, 0, 300, 750))
    im.save(half)
    plan = {"brand": "exemplo", "slides": [
        {"layout": "capa_manchete", "photo": "Fotos/half.jpg", "headline": "Tecido\nimporta?"}]}
    _, warnings = render(plan, brand, tmp_path / "o.pptx", tmp_path / "cache")
    assert any("fundo misto" in w for w in warnings)


def test_lint_flags_filled_shape_and_foreign_font(brand, tmp_path):
    from renderer.primitives import TextStyle, blank_slide, new_presentation, rect, text_box
    prs = new_presentation(1080, 1350)
    slide = blank_slide(prs)
    rect(slide, "pilula", (0, 0, 100, 40), fill="#3F6C75", radius=20)
    text_box(slide, "t", (0, 0, 100, 40), ["oi"], TextStyle("Poppins SemiBold", 20, 24))
    path = tmp_path / "bad.pptx"
    prs.save(str(path))
    problems = lint(path, brand)
    assert any("preenchida" in p for p in problems)
    assert any("SemiBold" in p for p in problems)
    assert any("fora da tipografia" in p for p in problems)


def test_example_plan_is_valid():
    plan = json.loads((BRANDS / "exemplo/plans/exemplo-carrossel-caimento.json").read_text(encoding="utf-8"))
    brand = load_brand(BRANDS, "exemplo", Path("/nao-importa"))
    rep = validate(plan, brand)
    assert [e for e in rep.errors if "foto não encontrada" not in e] == []
