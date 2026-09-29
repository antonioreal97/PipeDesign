"""Rodada 2: variações do PPTX para acertar peso da fonte e entrelinha no Canva.

C2: famílias base ("Playfair Display", "Montserrat") + negrito ligado + entrelinha 90%
C3: famílias com peso no nome ("... SemiBold") + entrelinha exata em pontos
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

ROOT = Path(__file__).parent
OUT = ROOT / "out2"; OUT.mkdir(exist_ok=True)
PHOTO = ROOT / "work" / "foto-teste.jpg"
W, H = 1080, 1350
PX = 9525
px = lambda v: Emu(int(v * PX))
pt = lambda v: Pt(v * 0.75)
ACCENT = (200, 164, 110)
COPY = {"brand": "MARCA A", "headline": ["Nova coleção", "Inverno 26"], "sub": "Peças atemporais para os dias frios", "cta": "Conheça agora"}

def set_line_spacing(p, mode, value):
    pPr = p._p.get_or_add_pPr()
    for old in pPr.findall(qn("a:lnSpc")):
        pPr.remove(old)
    ln = pPr.makeelement(qn("a:lnSpc"), {})
    if mode == "pct":
        child = ln.makeelement(qn("a:spcPct"), {"val": str(int(value * 1000))})
    else:  # pontos
        child = ln.makeelement(qn("a:spcPts"), {"val": str(int(value * 100))})
    ln.append(child)
    pPr.insert(0, ln)

def build(variant, dst):
    fonts = {
        "C2": {"head": ("Playfair Display", True), "semi": ("Montserrat", True), "reg": ("Montserrat", False), "ln": ("pct", 90)},
        "C3": {"head": ("Playfair Display SemiBold", False), "semi": ("Montserrat SemiBold", False), "reg": ("Montserrat", False), "ln": ("pts", None)},
    }[variant]
    prs = Presentation(); prs.slide_width, prs.slide_height = px(W), px(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(PHOTO), 0, 0, px(W), px(H)).name = "FOTO"
    shade = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, px(H - 620), px(W), px(620)); shade.name = "SOMBRA"
    shade.fill.solid(); shade.fill.fore_color.rgb = RGBColor(12, 10, 12)
    clr = shade.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr")); clr.append(clr.makeelement(qn("a:alpha"), {"val": "62000"}))
    shade.line.fill.background()
    mark = s.shapes.add_shape(MSO_SHAPE.OVAL, px(80), px(80), px(44), px(44)); mark.name = "LOGO_MARCA"
    mark.fill.solid(); mark.fill.fore_color.rgb = RGBColor(*ACCENT); mark.line.fill.background()

    def text(name, left, top, width, height, lines, font, size, color, line_px, align=PP_ALIGN.LEFT, spacing=None, anchor=MSO_ANCHOR.TOP):
        family, bold = font
        tb = s.shapes.add_textbox(px(left), px(top), px(width), px(height)); tb.name = name
        tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            mode, val = fonts["ln"]
            set_line_spacing(p, mode, val if mode == "pct" else line_px * 0.75)
            r = p.add_run(); r.text = line
            f = r.font; f.name = family; f.size = pt(size); f.bold = bold; f.color.rgb = RGBColor(*color)
            if spacing is not None:
                r._r.get_or_add_rPr().set("spc", str(int(spacing * 0.75 * 100)))
        return tb

    text("LOGO", 140, 84, 400, 40, [COPY["brand"]], fonts["semi"], 30, (255, 255, 255), 36, spacing=6)
    text("HEADLINE", 80, 830, 900, 240, COPY["headline"], fonts["head"], 112, (255, 255, 255), 118)
    text("SUBHEADLINE", 80, 1080, 900, 60, [COPY["sub"]], fonts["reg"], 38, (238, 238, 238), 48)
    pill = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(80), px(1170), px(340), px(84)); pill.name = "CTA_FUNDO"
    pill.adjustments[0] = 0.5; pill.fill.solid(); pill.fill.fore_color.rgb = RGBColor(*ACCENT); pill.line.fill.background()
    text("CTA", 80, 1170, 340, 84, [COPY["cta"]], fonts["semi"], 30, (26, 26, 26), 36, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    prs.save(str(dst))

build("C2", OUT / "teste-C2-pptx-negrito-90.pptx")
build("C3", OUT / "teste-C3-pptx-semibold-exato.pptx")
for f in sorted(OUT.iterdir()): print(f.name, f.stat().st_size)
